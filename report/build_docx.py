# -*- coding: utf-8 -*-
"""Dung baocao.docx tu cac khoi noi dung da bien dich.

Chay: python report/build_docx.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))

from report_data import (CONG_NGHE, KIEM_THU, LOI, LOI_KIEM_CHUNG, LOI_TAI_LIEU,
                         PROJECT, RUI_RO, TAI_LIEU_THAM_KHAO, VIET_TAT, YEU_CAU)
from content_front import LOI_MOC_DAU
from content_ch1 import CHUONG_1
from content_ch1b import CHUONG_1B
from content_ch2 import CHUONG_2
from content_ch3 import CHUONG_3, CHUONG_4

FIG = "report/figures"
MAX_IMG_W = 15.5
MAX_IMG_H = 20.0

CAP_RE = re.compile(r"^(Hình|Bảng)\s+(\d+|[A-E])\.(\d+)\.\s*(.+)$")
seq_counters = {}
caption_log = []


def set_cell_border(cell, sz=4, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def add_field(paragraph, instr, result="[Cập nhật trường: nhấn F9]",
              dirty=True, size=11, italic=True):
    r = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    if dirty:
        fld.set(qn("w:dirty"), "true")
    r._r.append(fld)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = paragraph.add_run(result)
    set_font(r4, size=size, italic=italic)
    r5 = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def style_document(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)

    for name, size in (("Heading 1", 16), ("Heading 2", 14),
                       ("Heading 3", 13), ("Heading 4", 12)):
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 1.5
        s.paragraph_format.keep_with_next = True

    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(3.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.0)


def add_page_number_footer(doc):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Ket qua luu cua truong PAGE la "1", khong phai huong dan nhap F9:
        # Word tinh lai khi lap trang, nhung neu tep duoc mo bang trinh xem
        # khong cap nhat truong thi chi con hien so trang, khong hien huong dan.
        add_field(p, " PAGE ", result="1", size=12, italic=False)


def caption_paragraph(doc, label, chapter, number, text, keep_next=False):
    """Sinh caption co truong SEQ de Word tu tao danh muc hinh/bang.

    So thu tu duoc kiem tra khi dung: neu so tu bien dich lech so
    dem thi dung ban, thay vi tao bao cao co so hinh sai.
    """
    key = (label, chapter)
    seq_counters[key] = seq_counters.get(key, 0) + 1
    n = seq_counters[key]
    if n != number:
        raise SystemExit(
            f"LOI DANH SO: '{label} {chapter}.{number}. {text[:40]}' "
            f"nhung thu tu dung la {n}. Kiem lai tap noi dung.")

    p = doc.add_paragraph()
    try:
        p.style = doc.styles["Caption"]
    except KeyError:
        pass
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = keep_next

    r = p.add_run(f"{label} {chapter}.")
    set_font(r, size=11, bold=True)
    instr = f" SEQ {label} \\* ARABIC" + (" \\r 1" if n == 1 else "")
    add_field(p, instr, result=str(n), size=11, italic=False)
    r2 = p.add_run(f". {text}")
    set_font(r2, size=11, bold=True)
    caption_log.append(f"{label} {chapter}.{n}. {text}")
    return p


def make_caption(doc, raw, keep_next=False):
    m = CAP_RE.match(raw.strip())
    if not m:
        raise SystemExit(f"LOI CAPTION: khong khop mau '{raw[:70]}'")
    label, chapter, num, text = m.group(1), m.group(2), int(m.group(3)), m.group(4)
    return caption_paragraph(doc, label, chapter, num, text, keep_next)


def all_defects():
    """Tat ca cac muc loi da xu ly, kem nhom de phan biet."""
    rows = [(m, "Mã nguồn và cấu hình", *rest) for m, *rest in LOI]
    rows += [(m, "Tài liệu", *rest) for m, *rest in LOI_TAI_LIEU]
    rows += [(m, "Phát hiện khi kiểm chứng cuối", *rest) for m, *rest in LOI_KIEM_CHUNG]
    return rows


def kiem_thu_rows():
    rows = [[a, b, str(c), str(d), "Đạt"] for a, b, c, d in KIEM_THU]
    total = sum(r[2] for r in KIEM_THU)
    rows.append(["Tổng cộng", "5 bộ kiểm thử", str(total), "0", "Đạt"])
    return rows


def resolve_rows(cap, header, rows, widths):
    """Nap du lieu cho bang bi bo trong khi bien dich.

    Dinh danh theo noi dung caption chu khong phai theo so cot: bao 2.3
    va bao so rui ro deu co 7 cot nen so cot khong phan biet duoc.
    """
    if rows:
        return header, rows, widths
    if "rủi ro" in cap:
        return header, [list(r) for r in RUI_RO], widths
    if "mục lỗi" in cap or "lỗi và cách khắc phục" in cap:
        return header, all_defects(), widths
    if "kiểm thử" in cap and len(header or []) == 5:
        return header, kiem_thu_rows(), widths
    raise SystemExit(f"KHONG NAP DUOC DU LIEU cho bang: {cap[:60]}")


def render_table(doc, caption, header, rows, widths):
    if caption:
        make_caption(doc, caption, keep_next=True)

    ncols = len(header) if header else len(widths)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    if header:
        cells = t.add_row().cells
        for i, h in enumerate(header):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_after = Pt(2)
            run = para.add_run(str(h))
            set_font(run, size=10, bold=True)
            set_cell_border(cells[i])
        repeat_header(t.rows[0])

    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row[:ncols]):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            para.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i else
                              WD_ALIGN_PARAGRAPH.CENTER)
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_after = Pt(2)
            run = para.add_run(str(val))
            set_font(run, size=10)

    for r_ in t.rows:
        for i, c in enumerate(r_.cells):
            if i < len(widths):
                c.width = Cm(widths[i])
            set_cell_border(c)
    return t


def render_figure(doc, path, caption):
    from PIL import Image
    fp = Path(path)
    if not fp.exists():
        print(f"  [thieu hinh] {path}", file=sys.stderr)
        return
    with Image.open(fp) as im:
        w, h = im.size
    scale = min(MAX_IMG_W / w, MAX_IMG_H / h)
    width, height = w * scale, h * scale

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(fp), width=Cm(width), height=Cm(height))

    make_caption(doc, caption)


def render_blocks(doc, blocks, section_after_h1=False):
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            doc.add_paragraph(b[1], style="Heading 1")
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        if kind in ("h2", "h3", "h4"):
            doc.add_paragraph(b[1], style="Heading " + kind[1])
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.0)
            r = p.add_run(b[1])
            set_font(r)
        elif kind == "bul":
            for it in b[1]:
                p = doc.add_paragraph(style="List Bullet")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r = p.add_run(it)
                set_font(r)
        elif kind == "num":
            for it in b[1]:
                p = doc.add_paragraph(style="List Number")
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r = p.add_run(it)
                set_font(r)
        elif kind == "note":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            r = p.add_run(b[1])
            set_font(r, size=11, italic=True)
        elif kind == "fig":
            render_figure(doc, b[1], b[2])
        elif kind == "tbl":
            cap, header, rows, widths = b[1], b[2], b[3], b[4]
            header, rows, widths = resolve_rows(str(cap), header, rows, widths)
            render_table(doc, cap, header, rows, widths)


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO MÔN QUẢN LÝ DỰ ÁN PHẦN MỀM")
    set_font(r, size=18, bold=True)
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PROJECT["ten"])
    set_font(r, size=16, bold=True)
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Mã nguồn: {PROJECT['ten_khoa']}")
    set_font(r, size=13)
    doc.add_paragraph()
    doc.add_paragraph()

    fields = [
        ("Người thực hiện", "Hoàng Tuấn Kiệt"),
        ("Mã số sinh viên", "[CẦN BỔ SUNG]"),
        ("Lớp", "[CẦN BỔ SUNG]"),
        ("Giảng viên hướng dẫn", "[CẦN BỔ SUNG]"),
        ("Trường, khoa", "[CẦN BỔ SUNG]"),
        ("Niên khóa", "[CẦN BỔ SUNG]"),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in fields:
        cells = t.add_row().cells
        cells[0].text = ""
        cells[1].text = ""
        r0 = cells[0].paragraphs[0].add_run(k)
        r0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(r0, size=12, bold=True)
        r1 = cells[1].paragraphs[0].add_run(v)
        set_font(r1, size=12)
        cells[0].width = Cm(5.0)
        cells[1].width = Cm(8.0)

    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Tháng 10 năm 2026")
    set_font(r, size=13, bold=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_front_matter(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC")
    set_font(r, size=16, bold=True)
    doc.add_paragraph()
    pt = doc.add_paragraph()
    add_field(pt, ' TOC \\o "1-4" \\h \\z \\u ')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DANH MỤC HÌNH")
    set_font(r, size=16, bold=True)
    doc.add_paragraph()
    pf = doc.add_paragraph()
    add_field(pf, ' TOC \\h \\z \\c "Hình" ')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DANH MỤC BẢNG")
    set_font(r, size=16, bold=True)
    doc.add_paragraph()
    pt2 = doc.add_paragraph()
    add_field(pt2, ' TOC \\h \\z \\c "Bảng" ')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DANH MỤC TỪ VIẾT TẮT")
    set_font(r, size=16, bold=True)
    doc.add_paragraph()
    t = doc.add_table(rows=0, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.add_row().cells
    for i, h in enumerate(["STT", "Viết tắt", "Ý nghĩa"]):
        hdr[i].text = ""
        rr = hdr[i].paragraphs[0].add_run(h)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(rr, size=10, bold=True)
        set_cell_border(hdr[i])
    repeat_header(t.rows[0])
    for i, (ab, mean) in enumerate(VIET_TAT, 1):
        c = t.add_row().cells
        for j, val in enumerate([str(i), ab, mean]):
            c[j].text = ""
            c[j].paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.CENTER if j < 2
                                            else WD_ALIGN_PARAGRAPH.LEFT)
            rr = c[j].paragraphs[0].add_run(val)
            set_font(rr, size=10)
            set_cell_border(c[j])
    for r_ in t.rows:
        r_.cells[0].width = Cm(1.5)
        r_.cells[1].width = Cm(3.0)
        r_.cells[2].width = Cm(10.5)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def section_heading(doc, text):
    """Tieu de muc cap 1, canh giua, van duoc dua vao muc luc."""
    p = doc.add_paragraph(text, style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_references(doc):
    section_heading(doc, "TÀI LIỆU THAM KHẢO")
    for i, src in enumerate(TAI_LIEU_THAM_KHAO, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(f"[{i}] {src}")
        set_font(r, size=11)


def add_appendices(doc):
    section_heading(doc, "PHỤ LỤC")

    doc.add_paragraph("Phụ lục A. Ma trận truy vết yêu cầu", style="Heading 1")
    doc.add_paragraph(
        "Ma trận truy vết đầy đủ gồm 164 dòng được lưu tại tệp "
        "docs/project-management/REQUIREMENT_TRACEABILITY.md trong kho mã nguồn. "
        "Tóm tắt được trình bày tại bảng A.1.")
    render_table(doc, "Bảng A.1. Tóm tắt ma trận truy vết yêu cầu",
                 ["Trạng thái", "Số dòng", "Tỷ lệ", "Ý nghĩa"],
                 [["Đã xác minh", "148", "90,2%", "Có kiểm thử đã chạy"],
                  ["Xác minh một phần", "8", "4,9%", "Đã triển khai, chưa phủ kiểm thử"],
                  ["Khoảng trống", "7", "4,3%", "Chưa thực hiện"],
                  ["Ngoài phạm vi", "1", "0,6%", "Cố ý loại, có lý do"],
                  ["Tổng", "164", "100%", "Toàn bộ dòng yêu cầu trong ma trận truy vết"]],
                 [3.5, 2.0, 2.0, 7.5])

    doc.add_paragraph("Phụ lục B. Sổ rủi ro", style="Heading 1")
    doc.add_paragraph(
        "Sổ rủi ro đầy đủ gồm 18 mục được lưu tại tệp "
        "docs/project-management/RISK_AUDIT.md. Toàn bộ mười tám mục được trình bày tại bảng B.1.")
    render_table(doc, "Bảng B.1. Sổ rủi ro của dự án",
                 ["Mã", "Rủi ro", "K", "H", "Điểm", "Trạng thái", "Kiểm soát chính"],
                 [list(r_) for r_ in RUI_RO],
                 [1.2, 3.6, 0.7, 0.7, 1.0, 2.2, 5.6])

    doc.add_paragraph("Phụ lục C. Bằng chứng kiểm thử", style="Heading 1")
    doc.add_paragraph(
        "Tệp kết quả chi tiết được lưu tại thư mục tests/results/ với dấu thời gian "
        "của lần chạy. Kết quả tổng hợp được trình bày tại bảng C.1.")
    render_table(doc, "Bảng C.1. Tổng hợp bằng chứng kiểm thử",
                 ["Bộ kiểm thử", "Lệnh", "Kiểm tra", "Lỗi", "Kết quả"],
                 kiem_thu_rows(),
                 [4.0, 5.2, 2.2, 1.8, 2.0])

    doc.add_paragraph("Phụ lục D. Danh mục sản phẩm bàn giao", style="Heading 1")
    doc.add_paragraph("Danh mục đầy đủ kèm cách xác minh được trình bày tại bảng D.1.")
    render_table(doc, "Bảng D.1. Danh mục sản phẩm bàn giao",
                 ["Sản phẩm", "Nội dung", "Cách xác minh"],
                 [["Mã nguồn backend", "78 tệp Java", "Đếm tệp trong backend/src/main/java"],
                  ["Ứng dụng di động", "22 tệp Kotlin", "Đếm tệp trong android/app/src/main/java"],
                  ["Bảng quản trị", "13 tệp TypeScript", "Đếm tệp trong admin-web/src"],
                  ["Cơ sở dữ liệu", "9 migration, 11 bảng", "Truy vấn information_schema"],
                  ["Triển khai", "3 hình ảnh, 3 container", "docker compose ps"],
                  ["Kiểm thử", "5 bộ, 385 kiểm tra", "Nhật ký kiểm thử"],
                  ["Tài liệu", "28 tệp Markdown", "Đếm tệp trong docs/"],
                  ["Thẻ phát hành", PROJECT["the"], "git tag"]],
                 [3.5, 4.5, 7.0])

    doc.add_paragraph("Phụ lục E. Nghiệm thu cuối cùng", style="Heading 1")
    doc.add_paragraph("Kết quả nghiệm thu từng hạng mục kèm bằng chứng được trình bày tại bảng E.1. Hàng cuối cùng cố ý không đánh dấu đạt, vì mức độ sẵn sàng vận hành chưa đạt.")
    render_table(doc, "Bảng E.1. Kết quả nghiệm thu cuối cùng",
                 ["Hạng mục", "Bằng chứng", "Trạng thái"],
                 [["Yêu cầu", "164 dòng: 148 xác minh, 8 một phần,\n7 khoảng trống, 1 ngoài phạm vi", "Đạt"],
                  ["Backend", "mvn verify: 113 kiểm thu, 0 lỗi", "Đạt"],
                  ["Ứng dụng di động", "10 kiểm thu, 0 lỗi", "Đạt"],
                  ["Bảng quản trị", "102/102 hợp đồng, 0 lỗ hổng phụ thuộc", "Đạt"],
                  ["Cơ sở dữ liệu", "9 migration, 11 bảng", "Đạt"],
                  ["AI", "6 ý định phân tích, đường dẫn từ chối hoạt động", "Đạt"],
                  ["Triển khai", "3/3 container khỏe", "Đạt"],
                  ["Kiểm thử tổng", "385 kiểm tra, 0 lỗi", "Đạt"],
                  ["Bảo mật", "0 bí mật trong kho mã nguồn", "Đạt"],
                  ["Mức độ sẵn sàng vận hành", "4 khoảng trống còn mở", "Không tuyên bố"]],
                 [3.8, 8.0, 3.2])


def main():
    doc = Document()
    style_document(doc)
    add_page_number_footer(doc)

    add_cover(doc)
    add_front_matter(doc)
    render_blocks(doc, LOI_MOC_DAU)
    render_blocks(doc, CHUONG_1)
    render_blocks(doc, CHUONG_1B)
    render_blocks(doc, CHUONG_2)
    render_blocks(doc, CHUONG_3)
    render_blocks(doc, CHUONG_4)
    add_references(doc)
    add_appendices(doc)

    # File dich bat buoc nam o thu muc goc cua du an.
    out = Path("baocao.docx")
    doc.save(out)
    log = Path("report/validation/captions.txt")
    log.parent.mkdir(parents=True, exist_ok=True)
    log.write_text("\n".join(caption_log), encoding="utf-8")
    print(f"Da tao: {out} ({out.stat().st_size:,} byte)")
    print(f"So caption: {len(caption_log)}")
    return out


if __name__ == "__main__":
    main()
